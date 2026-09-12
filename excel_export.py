from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from database import get_search, get_search_versions, get_batches, get_candidates_for_search

VOOLKIA_DARK = "240300"
VOOLKIA_ORANGE = "FF7000"
VOOLKIA_ORANGE_LIGHT = "FF9100"
VOOLKIA_GRAY = "EDEDED"
VOOLKIA_WHITE = "FFFFFF"


def _style_header(ws, row=1):
    fill = PatternFill("solid", fgColor=VOOLKIA_DARK)
    font = Font(color=VOOLKIA_WHITE, bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_width(ws, max_width=55):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        values = [str(c.value or "") for c in col]
        width = min(max(max(len(v) for v in values) + 2, 12), max_width)
        ws.column_dimensions[letter].width = width


def build_search_excel(search_id):
    search = get_search(search_id)
    versions = get_search_versions(search_id)
    batches = get_batches(search_id)
    candidates = get_candidates_for_search(search_id)

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Resumen")
    ws.append(["Campo", "Valor"])
    _style_header(ws)
    current = versions[0] if versions else {}

    for r in [
        ["Búsqueda", search["title"]],
        ["Estado", search["status"]],
        ["Zona", search["location"]],
        ["Seniority", search["seniority"]],
        ["Años mínimos", search["min_years"]],
        ["Total candidatos únicos", len(candidates)],
        ["Total tandas", len(batches)],
        ["Versión actual JD", current.get("version_number", "")],
        ["Fecha creación", search["created_at"]],
        ["Última actualización", search["updated_at"]],
    ]:
        ws.append(r)

    ws.append([])
    ws.append(["JD actual", current.get("jd", "")])
    ws.append(["Excluyentes", "\n".join(current.get("must_have", []))])
    ws.append(["Deseables", "\n".join(current.get("nice_to_have", []))])

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _auto_width(ws)

    ws = wb.create_sheet("Candidatos")
    headers = [
        "Nombre", "Ubicación", "Cargo actual", "Empresa actual", "Match %", "Prioridad",
        "Tanda", "Versión JD", "Evidencia", "A validar", "LinkedIn", "Fecha"
    ]
    ws.append(headers)
    _style_header(ws)
    ws.freeze_panes = "A2"

    max_batch = max([x["batch_number"] for x in candidates], default=0)
    for c in candidates:
        ws.append([
            c["name"], c["location"], c["current_role"], c["current_company"],
            c["match_score"], c["priority"], c["batch_number"], c["version_number"],
            c["evidence"], c["validation"], c["linkedin_url"], c["created_at"],
        ])
        row_idx = ws.max_row
        fill_color = VOOLKIA_ORANGE_LIGHT if c["batch_number"] == max_batch else VOOLKIA_GRAY
        fill = PatternFill("solid", fgColor=fill_color)
        for cell in ws[row_idx]:
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _auto_width(ws)

    ws = wb.create_sheet("Historial criterios")
    ws.append(["Versión", "Fecha", "Qué cambió", "Excluyentes", "Deseables", "JD"])
    _style_header(ws)
    for v in sorted(versions, key=lambda x: x["version_number"]):
        ws.append([
            v["version_number"], v["created_at"], v["change_note"],
            "\n".join(v["must_have"]), "\n".join(v["nice_to_have"]), v["jd"],
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _auto_width(ws)

    ws = wb.create_sheet("Tandas")
    ws.append(["Tanda", "Fecha", "Criterio usado", "Fuente / estrategia", "Candidatos nuevos"])
    _style_header(ws)
    for b in sorted(batches, key=lambda x: x["batch_number"]):
        ws.append([b["batch_number"], b["created_at"], b["criteria_summary"], b["source_summary"], b["candidate_count"]])
    _auto_width(ws)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
